"""Plex ODBC + Project Registry helpers shared by main.py (the interactive
Streamlit app) and daily_plex_export.py (the scheduled morning export). Split
out from main.py because main.py executes its whole UI top-to-bottom as soon
as it's imported, so it can't be imported directly from a non-Streamlit
script — this module only defines functions/constants and has no top-level
Streamlit UI calls, so it's safe to import from either."""

import io
import re

import pandas as pd
import streamlit as st
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

GSHEET_SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]


def get_credentials(scopes: list[str]):
    """Service-account credentials for the given OAuth scopes, built from the
    same gcp_service_account secret used by every Sheets/Drive caller."""
    from google.oauth2.service_account import Credentials

    return Credentials.from_service_account_info(st.secrets["gcp_service_account"], scopes=scopes)


@st.cache_resource
def get_gsheet_client():
    import gspread

    return gspread.authorize(get_credentials(GSHEET_SCOPES))


def is_registry_configured() -> bool:
    try:
        return "gcp_service_account" in st.secrets
    except Exception:
        return False


def get_or_create_worksheet(title: str, headers: list[str], rows: int = 200):
    """Returns the named worksheet tab in the shared registry spreadsheet
    (st.secrets["po_registry_sheet_id"]), creating it with a header row if it
    doesn't exist yet."""
    import gspread

    client = get_gsheet_client()
    sheet = client.open_by_key(st.secrets["po_registry_sheet_id"])
    try:
        return sheet.worksheet(title)
    except gspread.WorksheetNotFound:
        worksheet = sheet.add_worksheet(title=title, rows=rows, cols=len(headers))
        worksheet.append_row(headers)
        return worksheet


PROJECT_REGISTRY_HEADERS = [
    "Customer",
    "Project",
    "Part Number",
    "Contingency/Management Reserve Used",
    "Expected Project End Date",
    "NOTES",
]


def get_project_registry_worksheet():
    return get_or_create_worksheet("Project Registry", PROJECT_REGISTRY_HEADERS, rows=200)


def load_project_registry() -> pd.DataFrame:
    """Reads the hand-maintained Customer/Project/Part Number/Contingency/
    Expected End Date/NOTES rows. Project Balance is computed separately, not
    stored here. Returns an empty frame if Google Sheets isn't configured."""
    try:
        worksheet = get_project_registry_worksheet()
        records = worksheet.get_all_records()
        if not records:
            return pd.DataFrame(columns=PROJECT_REGISTRY_HEADERS)
        return pd.DataFrame(records)
    except Exception:
        return pd.DataFrame(columns=PROJECT_REGISTRY_HEADERS)


def parse_part_numbers(value) -> list[str]:
    """'931, 932, 985, 988' -> ['931', '932', '985', '988'], de-duplicated."""
    parts = re.split(r"[,\s]+", str(value).strip())
    seen = []
    for part in parts:
        if part and part not in seen:
            seen.append(part)
    return seen


def is_odbc_configured() -> bool:
    try:
        return "odbc_dsn" in st.secrets
    except Exception:
        return False


def get_odbc_connection():
    """A fresh pyodbc connection to Plex's cloud ODBC reporting service, via a
    named Windows DSN (Data Source Administrator) that has the host, port, and
    data source saved in it. The DataDirect OpenAccess SDK driver behind this
    DSN does *not* persist the UID/PWD from the setup dialog, so both are
    passed explicitly here — PWD carries the IAM access token
    (authmethod=iam;accesstoken=...), UID the Plex employee/user ID. Opened
    fresh per query rather than cached, since it's only used a handful of
    times per run and a long-lived cached connection can go stale between
    reruns. Only works on a machine that has the Plex ODBC driver and this
    DSN installed (Windows-only, set up via ODBC Data Source Administrator) —
    not available on Streamlit Community Cloud or other Linux hosts."""
    import pyodbc

    conn_str = (
        f"DSN={st.secrets['odbc_dsn']};"
        f"UID={st.secrets['odbc_uid']};"
        f"PWD={st.secrets['odbc_pwd']};"
    )
    return pyodbc.connect(conn_str)


# Plex tooling job-cost query, adapted from the version provided for the
# DataDirect OpenAccess SQL engine behind Plex's cloud ODBC service, which
# doesn't support T-SQL scripting (DECLARE, OPTION hints) — only a single
# SELECT. It also has a driver bug where a `?` bound parameter silently
# matches zero rows once more than one JOIN is involved (confirmed against
# the live service: identical literal SQL returns correct rows, the same
# query with the Part No. filter as a bound parameter returns none). So
# fetch_plex_job inlines the filter as an escaped, allowlist-validated
# literal via {part_no_pattern} instead of a query parameter — the Job Group
# / Job Key filters from the original query are dropped entirely since they
# were always left at -1 (no-op) in practice.
PLEX_JOB_QUERY_TEMPLATE = """
SELECT
    AJG.Accounting_Job_Group,
    AJ.Accounting_Job_No,
    P.Part_No + '-' + P.Revision            AS [Okay PN],
    P.Name                                  AS [Tooling Line Item Description],
    J.Job_No                                AS [Tooling Job No.],

    ISNULL(SP.Price,0) *
    ISNULL(SR.Quantity,0)                   AS [Total Revenue],
    ISNULL(AR.Credit,0)                     AS [Invoiced Revenue],
    ISNULL(PO_Cost.Cost,0)                  AS [Vendor POs Cost],
    ISNULL(LC.Cost,0)                       AS [Labor Cost],
    ISNULL(PO_Cost.Cost,0) +
    ISNULL(LC.Cost,0)                       AS [Total Cost],
    (ISNULL(SP.Price,0) *
     ISNULL(SR.Quantity,0)) -
    (ISNULL(PO_Cost.Cost,0) +
     ISNULL(LC.Cost,0))                     AS [Profit or Loss]

FROM Part_v_Part P
JOIN Part_v_Part_Product_Type PT
    ON PT.PCN = P.Plexus_Customer_No
    AND PT.Product_Type_Key = P.Product_Type_Key
LEFT OUTER JOIN Part_v_Job J
    ON P.Plexus_Customer_No = J.PCN
    AND P.Part_Key = J.Part_Key
LEFT OUTER JOIN Accounting_v_Accounting_Job AJ
    ON J.PCN = AJ.PCN
    AND J.Accounting_Job_Key = AJ.Accounting_Job_Key
LEFT OUTER JOIN Accounting_v_Accounting_Job_Group AJG
    ON AJ.PCN = AJG.PCN
    AND AJ.Accounting_Job_Group_Key = AJG.Accounting_Job_Group_Key
LEFT OUTER JOIN Sales_v_PO_Line SPOL
    ON P.Plexus_Customer_No = SPOL.PCN
    AND P.Part_Key = SPOL.Part_Key
LEFT OUTER JOIN Sales_v_Price SP
    ON SPOL.PCN = SP.PCN
    AND SPOL.PO_Line_Key = SP.PO_Line_Key
LEFT OUTER JOIN (
    SELECT SPOL2.PCN, SPOL2.PO_Line_Key,
           SUM(SR2.Quantity) AS Quantity
    FROM Sales_v_PO_Line SPOL2
    JOIN Sales_v_Release SR2
        ON SPOL2.PCN = SR2.PCN
        AND SPOL2.PO_Line_Key = SR2.PO_Line_Key
    GROUP BY SPOL2.PCN, SPOL2.PO_Line_Key
) AS SR
    ON SPOL.PCN = SR.PCN
    AND SPOL.PO_Line_Key = SR.PO_Line_Key
LEFT OUTER JOIN (
    SELECT ARID.Plexus_Customer_No, ARID.Part_Key,
           SUM(ARID.Credit) AS Credit
    FROM Accounting_v_AR_Invoice_Dist ARID
    GROUP BY ARID.Plexus_Customer_No, ARID.Part_Key
) AS AR
    ON P.Plexus_Customer_No = AR.Plexus_Customer_No
    AND P.Part_Key = AR.Part_Key
LEFT OUTER JOIN (
    SELECT POL.Plexus_Customer_No, POL.For_Part_Key,
           SUM(POL.Unit_Price * POR2.Quantity) AS Cost
    FROM Purchasing_v_Line_Item POL
    LEFT OUTER JOIN Purchasing_v_Release POR2
        ON POL.Plexus_Customer_No = POR2.Plexus_Customer_No
        AND POL.Line_Item_Key = POR2.Line_Item_Key
    GROUP BY POL.Plexus_Customer_No, POL.For_Part_Key
) AS PO_Cost
    ON P.Plexus_Customer_No = PO_Cost.Plexus_Customer_No
    AND P.Part_Key = PO_Cost.For_Part_Key
LEFT OUTER JOIN (
    SELECT C.PCN, C.Part_Key,
           SUM(ROUND(C.Extended_Cost,2)) AS Cost
    FROM Common_v_Cost C
    WHERE C.Cost_Sub_Type_Key = 17924
    GROUP BY C.PCN, C.Part_Key
) AS LC
    ON P.Plexus_Customer_No = LC.PCN
    AND P.Part_Key = LC.Part_Key

WHERE PT.Product_Type IN (
    'Tooling','Inspection Device','NRE/Tooling',
    'Packaging','In-Development','Production',
    'Protype','Service'
)
AND P.Part_No LIKE '{part_no_pattern}'

ORDER BY
    AJG.Accounting_Job_Group,
    AJ.Accounting_Job_No,
    P.Part_No + '-' + P.Revision
"""

# Allowlist for Part No. filters — since the driver bug above rules out real
# parameter binding, this is the injection defense for the literal
# substitution into PLEX_JOB_QUERY_TEMPLATE instead. Matches the character
# set actually seen in Plex Part Nos (letters, digits, spaces, -_.).
PART_NO_FILTER_RE = re.compile(r"^[A-Za-z0-9 _.-]+$")


def fetch_plex_job(part_no_filter: str) -> pd.DataFrame:
    """Runs the Plex tooling job-cost query for a single Part No. filter (e.g.
    '924'), returning one DataFrame — the ODBC equivalent of one Drive file."""
    if not PART_NO_FILTER_RE.match(part_no_filter):
        raise ValueError(f"Invalid Part No. filter: {part_no_filter!r}")
    pattern = part_no_filter.replace("'", "''") + "%"
    query = PLEX_JOB_QUERY_TEMPLATE.format(part_no_pattern=pattern)
    conn = get_odbc_connection()
    try:
        return pd.read_sql(query, conn)
    finally:
        conn.close()


def load_part_groups() -> list[list[str]]:
    """Part No. filters to fetch from Plex, read from the Parts.xlsx list
    (st.secrets["parts_list_path"]). A row like '985/931/988/930/931M' is one
    family — each number is queried separately but merged into a single
    sheet, the same '/'-joined convention used for labels in main.py (see
    derive_sheet_label). Shared by main.py (interactive fetch) and
    daily_plex_export.py (scheduled export) so both group families the same way."""
    df = pd.read_excel(st.secrets["parts_list_path"])
    values = df.iloc[:, 0].dropna().astype(str).str.strip()
    return [v.split("/") for v in values if v]


def fetch_plex_groups(groups: list[list[str]]) -> dict[str, pd.DataFrame]:
    """Runs one Plex query per Part No. in each group, concatenating a
    multi-member group (a family) into a single sheet keyed by the
    '/'-joined label instead of one sheet per number."""
    sheets = {}
    for group in groups:
        dfs = [fetch_plex_job(part_no) for part_no in group]
        sheets["/".join(group)] = pd.concat(dfs, ignore_index=True) if len(dfs) > 1 else dfs[0]
    return sheets


# Data cleaning + report building — shared by main.py (interactive app) and
# daily_email_report.py (scheduled daily email), so both produce identical
# cleaned/reported output from the same raw Plex export files.

STATUS_COLORS = {"Profit": "#0ca30c", "Loss": "#d03b3b"}
STATUS_COLORS_HEX = {"Profit": "0CA30C", "Loss": "D03B3B"}  # no '#', for openpyxl

REQUIRED_COLUMNS = [
    "Tooling Line Item Description",
    "Okay PN",
    "Tooling Job No.",
    "Total Revenue",
    "Invoiced Revenue",
    "Vendor POs Cost",
    "Labor Cost",
    "Total Cost",
    "Profit or Loss",
]

SUM_COLS = ["Invoiced Revenue", "Vendor POs Cost", "Labor Cost", "Total Cost"]

NUMERIC_TOTAL_COLS = [
    "Total Revenue",
    "Invoiced Revenue",
    "Vendor POs Cost",
    "Labor Cost",
    "Total Cost",
    "Profit or Loss",
    "Total PO $",
    "Budget Left",
]

EXCLUDED_OKAY_PNS = {
    "4066M-01",
    "388M-01",
    "388-1P-03",
    "388-1P-02",
    "1/1/4243",
    "4066 CHG-01",
    "4066 P CHG",
    "4243 CHG-01",
    "4304-P CHG-01",
    "938 ASY CHG",
    "939 ASY CHG",
    "985 ASY CHG",
    "931 CHG",
    "988 ASY CHG",
    "930 CHG",
    "930M",
    "931M",
    "388-1P-01",
}


def normalize_pn(value) -> str:
    """Normalizes Okay PN for registry matching: trims whitespace, collapses
    internal whitespace, strips a trailing dash (Plex exports sometimes add
    one, e.g. '932 A PD-01-' vs '932 A PD-01'), and ignores case."""
    text = re.sub(r"\s+", " ", str(value).strip().upper())
    return text.rstrip("- ")


def is_junk_okay_pn(value) -> bool:
    """True for an Okay PN that's just a bare number/revision code with no
    real description — e.g. '931-01', '924-1 -01', or a number plus a
    standalone 'ASY' suffix like '938 ASY-01'. These rows get dropped during
    cleaning. A PN with real text after ASY (e.g. '4243 ASY A MC-01-01') is
    kept — only the first, whole-word 'ASY' is stripped before checking."""
    text = re.sub(r"(?<![A-Z])ASY(?![A-Z])", "", str(value).strip().upper(), count=1)
    return bool(re.fullmatch(r"[\d\s\-]*", text))


def consolidate_duplicates(df: pd.DataFrame) -> pd.DataFrame:
    """Merge rows that share the same Okay PN, following the same rules as the
    original script: sum Total Revenue only when it differs across the group,
    zero out the other cost/revenue columns on the extra rows (or sum them
    onto the first row when they differ), then drop the extra rows."""
    df_cleaned = df[REQUIRED_COLUMNS].copy()
    df_work = df_cleaned.copy()
    rows_to_drop = []

    for okay_pn in df_work["Okay PN"].unique():
        group_indices = df_work[df_work["Okay PN"] == okay_pn].index.tolist()
        if len(group_indices) <= 1:
            continue

        group = df_work.loc[group_indices]
        first_idx, other_indices = group_indices[0], group_indices[1:]

        if len(group["Total Revenue"].unique()) != 1:
            df_work.loc[first_idx, "Total Revenue"] = group["Total Revenue"].sum()

        for col in SUM_COLS:
            if len(group[col].unique()) == 1:
                df_work.loc[other_indices, col] = 0
            else:
                df_work.loc[first_idx, col] = group[col].sum()

        rows_to_drop.extend(other_indices)

    df_cleaned = df_work.drop(rows_to_drop).reset_index(drop=True)
    df_cleaned["Profit or Loss"] = df_cleaned["Total Revenue"] - df_cleaned["Total Cost"]
    return df_cleaned


def load_all_sheets(uploaded_files) -> dict[str, pd.DataFrame]:
    sheets = {}
    for uploaded_file in uploaded_files:
        name = uploaded_file.name
        if name.lower().endswith(".csv"):
            sheets[name.rsplit(".", 1)[0]] = pd.read_csv(uploaded_file)
        elif name.lower().endswith((".xlsx", ".xls")):
            excel_file = pd.ExcelFile(uploaded_file)
            for sheet in excel_file.sheet_names:
                sheets[f"{name}_{sheet}"] = pd.read_excel(excel_file, sheet_name=sheet)
    return sheets


def clean_sheets(raw_sheets: dict[str, pd.DataFrame]) -> tuple[dict[str, pd.DataFrame], dict[str, list[str]]]:
    """Runs every raw sheet through the same drop-junk/exclude/dedup pipeline
    main.py uses after loading: returns (cleaned sheets, {sheet: missing
    columns} for any sheet skipped for lacking REQUIRED_COLUMNS)."""
    processed, missing_report = {}, {}
    for sheet_name, df in raw_sheets.items():
        missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
        if missing:
            missing_report[sheet_name] = missing
            continue
        df = df[df["Okay PN"].notna() & (df["Okay PN"].astype(str).str.strip() != "")]
        df = df[~df["Okay PN"].apply(is_junk_okay_pn)]
        df = df[~df["Okay PN"].map(normalize_pn).isin(EXCLUDED_OKAY_PNS)]
        processed[sheet_name] = consolidate_duplicates(df)
    return processed, missing_report


def is_po_registry_configured() -> bool:
    """True when the local PO Registry file (st.secrets["po_registry_local_path"])
    is configured — the source of truth for Total PO $, independent of Google
    Sheets (which still backs Project Registry/History)."""
    try:
        return "po_registry_local_path" in st.secrets
    except Exception:
        return False


def load_po_registry() -> dict[str, float]:
    """Returns {Okay PN: Total PO $}, read from the 'PO Registry' tab of the
    local Excel file at st.secrets["po_registry_local_path"] — an exact local
    mirror of the (now retired) Google Sheet registry. Returns an empty
    registry (rather than raising) if the file isn't configured/reachable, so
    the app still works without it — see PO_REGISTRY_SETUP.md."""
    try:
        df = pd.read_excel(st.secrets["po_registry_local_path"], sheet_name="PO Registry")
    except Exception:
        return {}

    registry = {}
    for _, r in df.iterrows():
        pn = str(r.get("Okay PN", "")).strip()
        if not pn:
            continue
        try:
            registry[normalize_pn(pn)] = float(r["Total PO $"])
        except (KeyError, TypeError, ValueError):
            continue
    return registry


def is_local_export_configured() -> bool:
    """True when a local Plex export folder is configured (the same folder
    daily_plex_export.py writes to, Drive-synced to the Cloud deployment's
    source folder) — lets this machine read the already-downloaded files
    directly off disk instead of a live Plex query or a Drive API round-trip
    for a copy that's already sitting locally."""
    try:
        return "plex_export_local_dir" in st.secrets
    except Exception:
        return False


def fetch_local_files():
    """CSV/Excel files already written by daily_plex_export.py into the local
    Drive-synced export folder. Path objects are a drop-in for load_all_sheets
    (same .name attribute, pd.read_csv/pd.ExcelFile both accept a Path)."""
    from pathlib import Path

    folder = Path(st.secrets["plex_export_local_dir"])
    return sorted(p for p in folder.glob("*") if p.suffix.lower() in (".csv", ".xlsx", ".xls"))


def add_totals_row(df: pd.DataFrame) -> pd.DataFrame:
    totals = {col: (df[col].sum() if col in NUMERIC_TOTAL_COLS else "") for col in df.columns}
    totals[df.columns[0]] = "TOTAL"
    return pd.concat([df, pd.DataFrame([totals])], ignore_index=True)


def extract_pn_prefix(value: str) -> str:
    """Leading run of letters/digits, stopping at the first dash, dot, space,
    underscore, etc. — e.g. '924-1' and '924-01' both become '924'."""
    match = re.match(r"^[A-Za-z0-9]+", value)
    return match.group(0) if match else value


def derive_sheet_label(df: pd.DataFrame, fallback: str) -> str:
    """Human-readable heading for a sheet, e.g. '932/4066', based on its Okay
    PN prefixes — used instead of the raw uploaded filename, which is often an
    auto-generated export name like 'Query_2026_07_10-09-12-16'. Falls back to
    that filename if no prefixes can be derived (e.g. an empty sheet)."""
    data_rows = df[df[df.columns[0]] != "TOTAL"]
    prefixes = data_rows["Okay PN"].dropna().astype(str).map(extract_pn_prefix).unique()
    return "/".join(prefixes) if len(prefixes) else fallback


def build_project_summary(final_sheets: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """One row per uploaded file, reusing each file's existing TOTAL row from
    Step 3. If a file contains multiple unique Tooling Job No. (or Okay PN
    prefix) values, they are joined with '/'."""
    rows = []
    for sheet_name, df in final_sheets.items():
        data_rows = df[df[df.columns[0]] != "TOTAL"]
        totals_row = df.iloc[-1]
        job_numbers = data_rows["Tooling Job No."].dropna().astype(str).unique()
        pn_prefixes = data_rows["Okay PN"].dropna().astype(str).map(extract_pn_prefix).unique()

        rows.append(
            {
                "File": sheet_name,
                "Project": "/".join(pn_prefixes),
                "Tooling Job No.": "/".join(job_numbers),
                "Total PO $": totals_row["Total PO $"],
                "Total Cost": totals_row["Total Cost"],
            }
        )

    summary = pd.DataFrame(rows)
    summary["Profit or Loss ($)"] = summary["Total PO $"] - summary["Total Cost"]
    summary["Status"] = summary["Profit or Loss ($)"].apply(lambda v: "Profit" if v >= 0 else "Loss")
    return summary


def build_excel(final_sheets: dict[str, pd.DataFrame]) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for sheet_name, df_final in final_sheets.items():
            df_final.to_excel(writer, sheet_name=sheet_name[:31], index=False)
    return buffer.getvalue()


def _write_project_summary_sheet(writer, project_summary: pd.DataFrame) -> None:
    """Writes the 'Project Summary' sheet (styled header, currency formatting,
    column widths, and a colored Profit/Loss bar chart) into an already-open
    ExcelWriter — shared by build_project_summary_excel (standalone download)
    and build_daily_report_excel (combined email attachment)."""
    project_summary.to_excel(writer, sheet_name="Project Summary", index=False)
    worksheet = writer.sheets["Project Summary"]
    n_rows = len(project_summary)

    header_fill = PatternFill("solid", fgColor="2A78D6")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    currency_cols = ["Total PO $", "Total Cost", "Profit or Loss ($)"]
    for col_name in currency_cols:
        col_letter = get_column_letter(project_summary.columns.get_loc(col_name) + 1)
        for row in range(2, n_rows + 2):
            worksheet[f"{col_letter}{row}"].number_format = '"$"#,##0.00'

    for i, col_name in enumerate(project_summary.columns, start=1):
        max_len = max(project_summary[col_name].astype(str).map(len).max(), len(col_name)) + 2
        worksheet.column_dimensions[get_column_letter(i)].width = max_len

    chart = BarChart()
    chart.type = "col"
    chart.title = "Profit or Loss by Project"
    chart.y_axis.title = "Profit or Loss ($)"
    chart.x_axis.title = "Project"
    chart.legend = None
    chart.height, chart.width = 10, 20

    profit_col = project_summary.columns.get_loc("Profit or Loss ($)") + 1
    project_col = project_summary.columns.get_loc("Project") + 1
    data = Reference(worksheet, min_col=profit_col, min_row=1, max_row=n_rows + 1)
    cats = Reference(worksheet, min_col=project_col, min_row=2, max_row=n_rows + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    series = chart.series[0]
    # Excel's default "invert if negative" recolors negative (Loss) bars with
    # its own automatic shade, silently overriding the dPt.solidFill below —
    # this is why the red never showed up. Must be off before setting colors.
    series.invertIfNegative = False
    series.data_points = [DataPoint(idx=i) for i in range(n_rows)]
    for i, status in enumerate(project_summary["Status"]):
        point = series.data_points[i]
        point.graphicalProperties.solidFill = STATUS_COLORS_HEX[status]
        point.graphicalProperties.line.noFill = True

    series.dLbls = DataLabelList()
    series.dLbls.showVal = True
    series.dLbls.numFmt = '"$"#,##0'

    worksheet.add_chart(chart, f"{get_column_letter(len(project_summary.columns) + 2)}2")


def build_project_summary_excel(project_summary: pd.DataFrame) -> bytes:
    """Project summary table plus a native, editable Excel bar chart colored
    green/red by Profit/Loss status."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        _write_project_summary_sheet(writer, project_summary)
    return buffer.getvalue()


def build_daily_report_excel(final_sheets: dict[str, pd.DataFrame], project_summary: pd.DataFrame) -> bytes:
    """Everything build_excel + build_project_summary_excel produce,
    combined into one workbook: one sheet per cleaned project plus a
    'Project Summary' sheet with its chart — used for the daily email
    attachment."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for sheet_name, df_final in final_sheets.items():
            df_final.to_excel(writer, sheet_name=sheet_name[:31], index=False)
        _write_project_summary_sheet(writer, project_summary)
    return buffer.getvalue()
